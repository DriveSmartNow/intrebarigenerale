#!/usr/bin/env python3
"""
Script pentru procesarea întrebărilor DRPCIV din Word

Extrage întrebări, răspunsuri, răspunsuri corecte și imagini din fișierul Word
și generează un JSON structurat pentru aplicația de învățare.
"""

import json
import os
import re
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from PIL import Image
from tqdm import tqdm


# Configurare logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('extract_questions.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)


class DrpcivQuestionExtractor:
    """Extractor pentru întrebările DRPCIV din fișiere Word."""
    
    def __init__(self, input_file: str, output_dir: str = "."):
        """
        Inițializează extractorul.
        
        Args:
            input_file: Calea către fișierul Word de procesat
            output_dir: Directorul pentru fișierele de ieșire
        """
        self.input_file = Path(input_file)
        self.output_dir = Path(output_dir)
        self.images_dir = self.output_dir / "imagini"
        self.questions: List[Dict[str, Any]] = []
        self.current_question_id = 1
        
        # Creează directoarele necesare
        self.images_dir.mkdir(exist_ok=True)
        
        # Pattern-uri pentru identificarea structurii
        self.question_patterns = [
            r'^\d+\.\s*(.+)',  # Întrebări numerotate: "1. Textul întrebării"
            r'^(\d+)\)\s*(.+)',  # Întrebări cu paranteză: "1) Textul întrebării"
            r'^Întrebarea\s+(\d+)[:\.]?\s*(.+)',  # "Întrebarea 1: Text"
        ]
        
        self.answer_pattern = r'^([ABC])\.\s*(.+)'  # Răspunsuri: "A. Text răspuns"
        self.correct_answer_patterns = [
            r'(răspuns|corect|correct|answer)[:\s]*([ABC])',
            r'([ABC])\s*[-–]\s*(corect|correct)',
            r'varianta\s+([ABC])',
            r'solutia\s+([ABC])',
        ]

    def extract_images_from_document(self, doc: DocumentType) -> Dict[str, str]:
        """
        Extrage imaginile din document și le salvează.
        
        Args:
            doc: Documentul Word
            
        Returns:
            Dict cu maparea între ID-urile imaginilor și căile salvate
        """
        image_map = {}
        image_counter = 1
        
        try:
            # Accesează partea de relații pentru imagini
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    try:
                        # Obține datele imaginii
                        image_data = rel.target_part.blob
                        
                        # Determină extensia fișierului
                        content_type = rel.target_part.content_type
                        if 'jpeg' in content_type or 'jpg' in content_type:
                            ext = 'jpg'
                        elif 'png' in content_type:
                            ext = 'png'
                        elif 'gif' in content_type:
                            ext = 'gif'
                        else:
                            ext = 'jpg'  # default
                        
                        # Salvează imaginea
                        image_filename = f"intrebare_{image_counter}.{ext}"
                        image_path = self.images_dir / image_filename
                        
                        with open(image_path, 'wb') as img_file:
                            img_file.write(image_data)
                        
                        # Optimizează imaginea pentru web
                        self._optimize_image(image_path)
                        
                        image_map[rel.rId] = f"imagini/{image_filename}"
                        image_counter += 1
                        
                        logger.info(f"Imaginea salvată: {image_filename}")
                        
                    except Exception as e:
                        logger.warning(f"Eroare la extragerea imaginii {rel.rId}: {e}")
                        
        except Exception as e:
            logger.error(f"Eroare la extragerea imaginilor: {e}")
            
        return image_map

    def _optimize_image(self, image_path: Path) -> None:
        """
        Optimizează o imagine pentru web.
        
        Args:
            image_path: Calea către imagine
        """
        try:
            with Image.open(image_path) as img:
                # Redimensionează dacă e prea mare
                max_size = (800, 600)
                if img.size[0] > max_size[0] or img.size[1] > max_size[1]:
                    img.thumbnail(max_size, Image.Resampling.LANCZOS)
                
                # Salvează cu compresie optimizată
                if image_path.suffix.lower() in ['.jpg', '.jpeg']:
                    img.save(image_path, 'JPEG', optimize=True, quality=85)
                elif image_path.suffix.lower() == '.png':
                    img.save(image_path, 'PNG', optimize=True)
                    
        except Exception as e:
            logger.warning(f"Eroare la optimizarea imaginii {image_path}: {e}")

    def parse_paragraph_text(self, paragraph) -> Tuple[str, Optional[str]]:
        """
        Parsează textul unui paragraf și identifică imaginile asociate.
        
        Args:
            paragraph: Paragraful de procesat
            
        Returns:
            Tuple cu textul și ID-ul imaginii (dacă există)
        """
        text = paragraph.text.strip()
        image_id = None
        
        # Verifică dacă paragraful conține imagini
        for run in paragraph.runs:
            try:
                # Caută elemente de tip blip (imagini)
                for drawing in run._element.iter():
                    if drawing.tag.endswith('}blip'):
                        embed_id = drawing.get(qn('r:embed'))
                        if embed_id:
                            image_id = embed_id
                            break
            except Exception:
                continue
        
        return text, image_id

    def identify_question_start(self, text: str) -> Optional[Dict[str, Any]]:
        """
        Identifică începutul unei întrebări noi.
        
        Args:
            text: Textul de analizat
            
        Returns:
            Dict cu informații despre întrebare dacă e identificată, None altfel
        """
        for pattern in self.question_patterns:
            match = re.match(pattern, text, re.IGNORECASE)
            if match:
                if len(match.groups()) == 1:
                    return {
                        'id': self.current_question_id,
                        'text': match.group(1).strip(),
                        'answers': [],
                        'correct_answer': None,
                        'image': None
                    }
                elif len(match.groups()) == 2:
                    return {
                        'id': self.current_question_id,
                        'text': match.group(2).strip(),
                        'answers': [],
                        'correct_answer': None,
                        'image': None
                    }
        return None

    def identify_answer(self, text: str) -> Optional[Dict[str, str]]:
        """
        Identifică un răspuns multiple choice.
        
        Args:
            text: Textul de analizat
            
        Returns:
            Dict cu litera și textul răspunsului dacă e identificat
        """
        match = re.match(self.answer_pattern, text, re.IGNORECASE)
        if match:
            return {
                'letter': match.group(1).upper(),
                'text': match.group(2).strip()
            }
        return None

    def identify_correct_answer(self, text: str) -> Optional[str]:
        """
        Identifică răspunsul corect.
        
        Args:
            text: Textul de analizat
            
        Returns:
            Litera răspunsului corect dacă e identificată
        """
        for pattern in self.correct_answer_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            if matches:
                for match in matches:
                    if isinstance(match, tuple):
                        for item in match:
                            if item.upper() in ['A', 'B', 'C']:
                                return item.upper()
                    elif isinstance(match, str) and match.upper() in ['A', 'B', 'C']:
                        return match.upper()
        return None

    def process_document(self) -> None:
        """Procesează documentul Word și extrage întrebările."""
        logger.info(f"Începe procesarea documentului: {self.input_file}")
        
        try:
            # Încarcă documentul
            doc = Document(self.input_file)
            logger.info(f"Document încărcat cu succes. Numărul de paragrafe: {len(doc.paragraphs)}")
            
            # Extrage imaginile
            logger.info("Extragerea imaginilor...")
            image_map = self.extract_images_from_document(doc)
            logger.info(f"Extrase {len(image_map)} imagini")
            
            # Procesează paragrafele
            current_question = None
            current_image_id = None
            
            for i, paragraph in enumerate(tqdm(doc.paragraphs, desc="Procesare paragrafe")):
                text, image_id = self.parse_paragraph_text(paragraph)
                
                if not text:
                    continue
                
                # Păstrează referința la imaginea curentă
                if image_id:
                    current_image_id = image_id
                
                # Verifică dacă e începutul unei întrebări noi
                question_info = self.identify_question_start(text)
                if question_info:
                    # Salvează întrebarea anterioară dacă există
                    if current_question and self._is_question_valid(current_question):
                        self.questions.append(current_question)
                    
                    # Începe o întrebare nouă
                    current_question = question_info
                    current_question['id'] = self.current_question_id
                    self.current_question_id += 1
                    
                    # Asociază imaginea dacă există
                    if current_image_id and current_image_id in image_map:
                        current_question['image'] = image_map[current_image_id]
                    
                    continue
                
                # Verifică dacă e un răspuns
                answer_info = self.identify_answer(text)
                if answer_info and current_question:
                    formatted_answer = f"{answer_info['letter']}) {answer_info['text']}"
                    current_question['answers'].append(formatted_answer)
                    continue
                
                # Verifică dacă e răspunsul corect
                correct_answer = self.identify_correct_answer(text)
                if correct_answer and current_question:
                    current_question['correct_answer'] = correct_answer
                    continue
                
                # Verifică dacă este o linie cu categorie și informații suplimentare
                if re.search(r'Categoria:.*\|.*Grad\s+dificultate', text, re.IGNORECASE):
                    # Finalizează întrebarea curentă
                    if current_question and self._is_question_valid(current_question):
                        self.questions.append(current_question)
                        current_question = None
                    continue
                
                # Verifică dacă este o linie cu comentarii
                if re.search(r'\d+\s+comentarii', text):
                    continue
                
                # Dacă nu se potrivește cu niciun pattern, poate fi continuarea întrebării
                if current_question and not current_question['answers']:
                    current_question['text'] += " " + text
            
            # Salvează ultima întrebare
            if current_question and self._is_question_valid(current_question):
                self.questions.append(current_question)
            
            logger.info(f"Procesare completă. Extrase {len(self.questions)} întrebări")
            
        except Exception as e:
            logger.error(f"Eroare la procesarea documentului: {e}")
            raise

    def _is_question_valid(self, question: Dict[str, Any]) -> bool:
        """
        Verifică dacă o întrebare este validă.
        
        Args:
            question: Întrebarea de validat
            
        Returns:
            True dacă întrebarea este validă
        """
        return (
            question.get('text') and
            len(question.get('answers', [])) >= 2  # Acceptă întrebări fără răspuns corect marcat
        )

    def generate_json_output(self) -> Dict[str, Any]:
        """
        Generează structura JSON finală.
        
        Returns:
            Dict cu structura JSON
        """
        # Formatează întrebările pentru output
        formatted_questions = []
        for q in self.questions:
            formatted_q = {
                'id': q['id'],
                'intrebare': q['text'],
                'raspunsuri': q['answers']
            }
            
            # Adaugă răspunsul corect doar dacă există
            if q.get('correct_answer'):
                formatted_q['raspuns_corect'] = q['correct_answer']
            
            if q.get('image'):
                formatted_q['imagine'] = q['image']
            
            # Adaugă explicație dacă există
            if q.get('explanation'):
                formatted_q['explicatie'] = q['explanation']
            
            formatted_questions.append(formatted_q)
        
        # Creează structura finală
        output_data = {
            'categorii': [
                {
                    'nume': 'Mediu de învățare general',
                    'descriere': 'Întrebări generale pentru examenul auto DRPCIV',
                    'intrebari': formatted_questions
                }
            ],
            'total_intrebari': len(formatted_questions),
            'versiune': '1.0',
            'ultima_actualizare': datetime.now().strftime('%Y-%m-%d')
        }
        
        return output_data

    def save_json_output(self, output_data: Dict[str, Any]) -> None:
        """
        Salvează datele în format JSON.
        
        Args:
            output_data: Datele de salvat
        """
        output_file = self.output_dir / 'intrebari-generale.json'
        
        try:
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(output_data, f, indent=2, ensure_ascii=False)
            
            logger.info(f"JSON salvat în: {output_file}")
            
        except Exception as e:
            logger.error(f"Eroare la salvarea JSON: {e}")
            raise

    def run(self) -> None:
        """Rulează întregul proces de extracție."""
        logger.info("=== Începe extracția întrebărilor DRPCIV ===")
        
        # Verifică dacă fișierul de intrare există
        if not self.input_file.exists():
            raise FileNotFoundError(f"Fișierul {self.input_file} nu a fost găsit")
        
        # Procesează documentul
        self.process_document()
        
        # Generează JSON-ul
        logger.info("Generarea output-ului JSON...")
        output_data = self.generate_json_output()
        
        # Salvează rezultatele
        self.save_json_output(output_data)
        
        # Raport final
        logger.info("=== Extracția completă ===")
        logger.info(f"Întrebări extrase: {len(self.questions)}")
        logger.info(f"Imagini salvate: {len(list(self.images_dir.glob('*')))}")
        logger.info(f"JSON salvat: intrebari-generale.json")


def main():
    """Funcția principală."""
    import argparse
    
    parser = argparse.ArgumentParser(description='Extrage întrebări DRPCIV din fișiere Word')
    parser.add_argument(
        'input_file',
        nargs='?',
        default='473514317-227353061-Intrebari-Drpciv-docx.docx',
        help='Fișierul Word de procesat'
    )
    parser.add_argument(
        '--output-dir',
        default='.',
        help='Directorul pentru fișierele de ieșire'
    )
    
    args = parser.parse_args()
    
    try:
        extractor = DrpcivQuestionExtractor(args.input_file, args.output_dir)
        extractor.run()
        
    except Exception as e:
        logger.error(f"Eroare în execuția scriptului: {e}")
        return 1
    
    return 0


if __name__ == '__main__':
    exit(main())